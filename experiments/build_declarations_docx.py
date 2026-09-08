#!/usr/bin/env python3
"""Build the declaration files Editorial Manager wants as separate .docx uploads.

Pattern Recognition's attach-files step asks for a competing-interests declaration as
its own editable document rather than only as a manuscript section, so this writes it
(and the author statement it sits with) from the same text the manuscript carries.

Run: python experiments/build_declarations_docx.py
"""
from __future__ import annotations

import pathlib

from docx import Document
from docx.shared import Pt

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript"

TITLE = ("WAGER: An Exact Decomposition of Model-Pair Score Gains on Benchmarks "
         "with Strong Label Priors")
AUTHOR = "Quang-Vinh Dang, British University Vietnam, Hung Yen, Vietnam"

BODY = [
    ("Declaration of competing interests", [
        "The author declares that he has no known competing financial interests or "
        "personal relationships that could have appeared to influence the work "
        "reported in this paper.",
    ]),
    ("Funding", [
        "This research did not receive any specific grant from funding agencies in "
        "the public, commercial, or not-for-profit sectors.",
    ]),
    ("CRediT authorship contribution statement", [
        "Quang-Vinh Dang: Conceptualization; Methodology; Formal analysis; Software; "
        "Investigation; Validation; Visualization; Writing – original draft; "
        "Writing – review and editing.",
    ]),
    ("Declaration of generative AI and AI-assisted technologies in the manuscript "
     "preparation process", [
        "During the preparation of this work the author used generative AI tools in "
        "order to support the writing of the manuscript and the development of the "
        "accompanying code. The author reviewed and edited the output as needed and "
        "takes full responsibility for the content of the published article.",
    ]),
]


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    doc.add_paragraph("Declarations").bold = True
    doc.add_paragraph(f"Manuscript: {TITLE}")
    doc.add_paragraph(f"Author: {AUTHOR}")
    doc.add_paragraph()

    for heading, paras in BODY:
        h = doc.add_paragraph()
        h.add_run(heading).bold = True
        for text in paras:
            doc.add_paragraph(text)
        doc.add_paragraph()

    dest = OUT / "declarations.docx"
    doc.save(dest)
    print(f"wrote {dest}")


if __name__ == "__main__":
    main()
